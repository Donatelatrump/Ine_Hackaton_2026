"""Verifica cifras del caso, documentos regenerados y composición de la entrega."""
from pathlib import Path
import json
import hashlib
import sys
from datetime import date
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / '.task_packages'))
sys.path.insert(0, str(Path(__file__).resolve().parent / 'reporte' / '.packages'))
sys.path.insert(0, str(Path(__file__).resolve().parent / 'presentacion' / '.packages'))
import pymupdf
from pptx import Presentation
from PIL import Image

HERE = Path(__file__).resolve().parent
TRACK = HERE.parent
case = json.loads((HERE/'caso_municipal.json').read_text(encoding='utf-8'))
data = json.loads((TRACK/'demo/data/resultados.json').read_text(encoding='utf-8'))
r = next(r for r in data['municipios'] if r['key'] == '23008')
comparison=json.loads((TRACK/'demo/data/comparacion_modelos.json').read_text(encoding='utf-8'))
forecasts={row['key']:row for row in comparison['forecasts']}
method=comparison['report']['selected_model']
assert case['configuration']['method']==method
def position(model):
    rows=[row for row in data['municipios'] if row['stable_catalog_state'] and forecasts.get(row['key'],{}).get(model) is not None]
    rows.sort(key=lambda row:(-forecasts[row['key']][model],row['key']))
    return next(i for i,row in enumerate(rows,1) if row['key']=='23008')
assert case['municipality']['observed_2024'] == r['latest_sv'] + r['latest_nv'] == 255362
assert case['position_model']==position(method)==9
assert case['position_ridge']==position('ridge')==4 and case['position_reference']==position('baseline')==3
assert case['stability']['count'] == case['stability']['eligible'] == 4
scenario = case['scenario']
assert abs(scenario['additional_hypothetical'] - 1000*case['evidence']['effect_pp']/100) < 1e-9
assert abs(scenario['before']-forecasts['23008'][method])<1e-9
assert abs(scenario['after'] - (forecasts['23008'][method] - 100*scenario['additional_hypothetical']/r['observed'])) < 1e-9
checks = []
for name, count in [('reporte/Umbral_reporte_metodologico_preventivo.pdf',14), ('presentacion/Umbral_presentacion.pdf',10)]:
    doc = pymupdf.open(HERE/name)
    assert len(doc) == count
    all_text = '\n'.join(p.get_text() for p in doc)
    assert '23008' in all_text and 'SMS' in all_text and '0.30' in all_text, name
    assert 'proporcional' in all_text.lower() and 'prioridad' in all_text.lower(), name
    assert 'estabilidad' in all_text.lower() and '2024' in all_text
    assert 'umbral' in all_text.lower() and 'demo7' not in all_text.lower() and '1.5' in all_text and 'semáforo' in all_text.lower()
    for index, p in enumerate(doc, 1):
        blocks = p.get_text('blocks')
        assert all(b[0] >= 0 and b[1] >= 0 and b[2] <= p.rect.width+1 and b[3] <= p.rect.height+1 for b in blocks), (name, index)
        ordered = sorted(blocks,key=lambda b:b[1])
        for a, b in zip(ordered,ordered[1:]):
            if a[0] < b[2] and b[0] < a[2]:
                assert a[3] <= b[1]+1, (name,index,'Texto superpuesto',a[:4],b[:4])
    checks.append({'file':name,'pages':count,'bounds_ok':True,'text_overlap':False,'sha256':hashlib.sha256((HERE/name).read_bytes()).hexdigest()})
prs = Presentation(HERE/'presentacion/Umbral_presentacion.pptx')
assert len(prs.slides) == 10
ppt_text = '\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
assert '23008' in ppt_text and '4/6' in ppt_text and 'SMS' in ppt_text
assert 'proporcional' in ppt_text.lower() and 'prioridad' in ppt_text.lower()
assert 'estabilidad' in ppt_text.lower() and 'cobertura' in ppt_text.lower()
checks.append({'file':'presentacion/Umbral_presentacion.pptx','slides':10,'case_and_decision_layers':True})
images = [Image.open(HERE/f'presentacion/QA_{i:02}.png').convert('RGB') for i in range(1,11)]
montage = Image.new('RGB',(1200,5*350),'#dfe8e3')
for i, img in enumerate(images):
    img.thumbnail((590,340))
    montage.paste(img,((i%2)*600+5,(i//2)*350+5))
montage.save(HERE/'presentacion/QA_montaje.png')
browser = json.loads((TRACK/'demo/verificacion/demo7/resultado.json').read_text(encoding='utf-8'))
assert browser['passed'] and not browser['page_errors']
for path,sha in browser['sources'].items():
    assert hashlib.sha256((TRACK/'demo'/path).read_bytes()).hexdigest()==sha,('Verificación web obsoleta',path)
result = {'date':date.today().isoformat(),'application':'Umbral','documents':checks,'case_reproduced':True,'browser_checks':len(browser['checks']),'browser_sources_verified':True,'scenario_action':case['configuration']['action'],'external_effect_is_hypothetical':True,'browser_errors':browser['page_errors'],'limitations':['PPTX no renderizado con PowerPoint','Modelo 2027 exploratorio; escenario sin efecto causal demostrado','Registro, elegibilidad, firmas y envío no acreditados']}
(HERE/'verificacion_unificacion.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('Entrega verificada: reporte 14 páginas, presentación 10, caso 23008 y geometría textual.')
