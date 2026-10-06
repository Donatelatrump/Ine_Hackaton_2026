"""Umbral: PPTX editable y PDF con la misma composición. Python 3 + python-pptx/reportlab.
Ejecutar desde esta carpeta. Dependencias locales opcionales en .packages.
"""
from pathlib import Path
import sys, textwrap, json
HERE=Path(__file__).resolve().parent
sys.path[:0]=[str(HERE/'.packages'),str(HERE.parent/'reporte'/'.packages')]
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

INK='15353B'; TEAL='10685F'; PAPER='F3F6F4'; MUTED='587177'; SAGE='E5EEEA'
CCPC='https://ine.mx/transparencia/datos-abiertos/visualizacion-datos/conteos-censales-participacion/'
CALL='https://ine.mx/wp-content/uploads/2026/09/Convocatoria_Hackaton-1.pdf'
slides=[
('UMBRAL', [('Atlas de decisión municipal',28),('Predicción exploratoria y escenarios de abstención para 2027',22),('Tecnocalidos · Hackathon INE 2026 · Track 2',18),('Martin Isai Nuñez Villeda y Cesar Donnet Hernandez Gastelum',17),('Versión técnica · 5 de octubre de 2026',17)],'Datos agregados CCPC 2009–2024. Prototipo para revisión institucional.'),
('Una capacidad de 50 municipios', [('La decisión inicial',24),('Un área de educación cívica puede revisar 50 municipios. Umbral ordena antecedentes, calidad y supuestos para preparar un plan de atención.',22),('La persona usuaria elige entidad, capacidad y criterio de prioridad. Revisa cada municipio antes de exportar la lista.',22),('50 es una capacidad propuesta y editable.',18)],'Fuente: convocatoria INE, pp. 3–5. '+CALL),
('Cobertura y calidad visibles', [('1,783 con estimación',36),('694 sin estimación · 2,477 claves históricas',24),('Abstención observada = 100 × NV / (SV + NV)',24),('NS identifica participación no especificada. La regla principal exige NS ≤10% en ambos antecedentes relevantes.',20),('El catálogo constante orienta la revisión territorial. Falta comprobar equivalencia de límites históricos.',20)],'Fuente: auditoría y base municipal del proyecto. CCPC: '+CCPC),
('La referencia también compite', [('Referencia: última elección del mismo tipo',24),('Para 2027 usa 2021. Ridge estima un ajuste con historia electoral, edad agregada, urbanización y calidad.',22),('Alpha = 10 permanece congelado. Cada corte aprende con resultados anteriores al año objetivo.',22),('2021 seleccionó el método. Ese corte no es una prueba independiente posterior a la selección.',20)],'Fuente: modelado/METODOLOGIA.md y modelado/salidas/metricas.json. No hay ajuste para mejorar 2024.'),
('Ridge mejora en 2021 y empeora en 2024', [],'MAE en puntos porcentuales. Ponderación por SV+NV observado del objetivo. Top50 dentro de cada cohorte evaluable.'),
('Mapa y escenario sobre los mismos territorios', [('Base, con intervención y diferencia',26),('La persona usuaria cambia alcance y respuesta supuesta. Los umbrales de la base permanecen fijos al comparar los mapas.',22),('La cobertura usa SV+NV observado en 2024. NS queda fuera del escenario y se informa por separado.',22),('Ejemplo hipotético: base 60%, 10,000 observados, alcance 1,000 y respuesta 5% producen 30 participantes adicionales y 59.7%.',20)],'Capturas de la demo corregida. El escenario es hipotético, sin efecto causal demostrado ni población 2027 proyectada.'),
('La prioridad depende del objetivo', [('Tasa',24),('Revisa municipios con mayor abstención estimada.',21),('Volumen sobre observados 2024',24),('Revisa dónde la tasa implica más abstenciones en la base observable.',21),('Deterioro respecto al mismo tipo electoral',24),('Revisa el aumento estimado frente al antecedente comparable.',21)],'Dos métodos × NS 5/10/20: frecuencia topN y concordancia describen estabilidad. No expresan confianza calibrada.'),
('Prevención con una ficha verificable', [('Cada municipio necesita una propuesta revisable',26),('La ficha relaciona antecedentes observables con una acción de información electoral accesible o educación cívica.',22),('Registra responsable propuesto, personas alcanzadas, condiciones de aplicación y medición. El mecanismo esperado es una hipótesis.',22),('Piloto propuesto: asignación aleatoria dentro de un conjunto elegible y balanceado alrededor del corte. Medir cobertura y participación.',20)],'La proximidad del score no demuestra equivalencia. El prototipo no acredita impacto causal ni un diseño RDD.'),
('Uso institucional con revisión humana', [('La vista por entidad resume capacidad y cobertura',25),('El área usuaria revisa calidad, concordancia de listas y viabilidad del alcance. La exportación conserva los supuestos de cada decisión.',22),('Las claves municipales permiten trazabilidad. Sin correspondencia verificada, el producto no atribuye municipios a distritos.',22),('La demo local y los datos agregados permiten inspección sin dependencias remotas.',21)],'Límites: pocas ocasiones electorales, cohortes filtradas y armonización territorial pendiente.'),
('Siguiente decisión: un piloto acotado', [('Validación con el área usuaria',26),('Acordar objetivo de priorización, capacidad real y protocolo de revisión territorial. Elegir acciones con cobertura que pueda medirse.',22),('Evaluación antes de ampliar',26),('Predefinir indicadores y asignación del piloto. Conservar la referencia y publicar resultados favorables y adversos.',22),('Equipo: Tecnocalidos. Firmas, registro y envío formal siguen pendientes.',19)],'2024 es presidencial y no valida directamente 2027. No se promete superioridad predictiva o efectividad preventiva.')]

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
fontpath=Path('C:/Windows/Fonts/segoeui.ttf'); boldpath=Path('C:/Windows/Fonts/segoeuib.ttf')
pdfmetrics.registerFont(TTFont('Segoe',str(fontpath))); pdfmetrics.registerFont(TTFont('SegoeBold',str(boldpath)))
pdf=canvas.Canvas(str(HERE/'Umbral_presentacion.pdf'),pagesize=(960,540))
def rgb(h): return RGBColor.from_string(h)
def text(slide,value,x,y,w,h,size=22,color=INK,bold=False):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name='Segoe UI'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
    pdf.setFillColor('#'+color); pdf.setFont('SegoeBold' if bold else 'Segoe',size)
    words=value.split(); lines=[]; current=''
    for word in words:
        test=(current+' '+word).strip()
        if pdfmetrics.stringWidth(test,'SegoeBold' if bold else 'Segoe',size)>w*72 and current: lines.append(current); current=word
        else: current=test
    if current: lines.append(current)
    box.height=Inches(max(h,len(lines)*size*1.22/72+.08))
    for i,line in enumerate(lines): pdf.drawString(x*72,540-y*72-size-i*size*1.22,line)
    return len(lines)*size*1.22/72

for num,(title,parts,foot) in enumerate(slides,1):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=rgb(PAPER)
    pdf.setFillColor('#'+PAPER); pdf.rect(0,0,960,540,fill=1,stroke=0)
    text(slide,'UMBRAL · '+str(num).zfill(2),.65,.35,12,.3,12,TEAL,True)
    text(slide,title,.65,.9,12,1.2,44 if num==1 else 32,TEAL if num==1 else INK,True)
    if num==5:
        data=[['Corte / método','n','MAE (pp)','Ponderado','Top50'],['2021 · Referencia','1,490','12.12','8.05','16/50'],['2021 · Ridge','1,490','9.21','5.49','19/50'],['2024 · Referencia','1,476','5.64','3.54','27/50'],['2024 · Ridge','1,476','8.58','4.49','27/50']]
        table=slide.shapes.add_table(5,5, Inches(.65), Inches(2.0), Inches(12), Inches(2.6)).table
        table.columns[0].width=Inches(4)
        for c in range(1,5): table.columns[c].width=Inches(2)
        for r,row in enumerate(data):
            for c,value in enumerate(row):
                cell=table.cell(r,c); cell.text=value; cell.fill.solid(); cell.fill.fore_color.rgb=rgb(SAGE if r==0 else PAPER)
                for p in cell.text_frame.paragraphs: p.font.name='Segoe UI'; p.font.size=Pt(20); p.font.bold=r==0; p.font.color.rgb=rgb(INK)
                text_x=.75+sum([4,2,2,2,2][:c])
                pdf.setFillColor('#'+INK); pdf.setFont('SegoeBold' if r==0 else 'Segoe',20); pdf.drawString(text_x*72,540-(2.1+r*.52)*72-20,value)
        text(slide,'La evidencia no demuestra superioridad consistente.',.65,4.95,12,.5,25,TEAL,True)
        text(slide,'Spearman referencia / Ridge: 0.485 / 0.599 en 2021 y 0.827 / 0.811 en 2024. Análisis posterior, sin reajuste del modelo.',.65,5.6,12,.7,20)
    elif num==6:
        text(slide,'Con intervención hipotética',.65,2.0,5.9,.4,22,TEAL,True)
        text(slide,'Diferencia respecto a la base',6.85,2.0,5.8,.4,22,TEAL,True)
        for filename,x in [('mapa_intervencion.png',.65),('mapa_cambio.png',6.85)]:
            img=HERE/filename
            slide.shapes.add_picture(str(img),Inches(x),Inches(2.55),width=Inches(5.85))
            pdf.drawImage(str(img),x*72,540-2.55*72-5.85*72*648/1026,width=5.85*72,height=5.85*72*648/1026)
        text(slide,'Mismos territorios y umbrales base fijos. Cobertura y ponderación: SV+NV observado 2024. NS se informa por separado.',.65,6.25,12,.5,16)
    else:
        y=2.15 if num!=1 else 2.6
        for value,size in parts:
            h=text(slide,value,.65,y,11.9,1.0 if size<28 else .7,size,TEAL if size>=24 else INK,size>=24)
            y+=h+.22
    text(slide,foot,.65,6.83,11.9,.45,10,MUTED)
    slide.notes_slide.notes_text_frame.text=foot+'\nFuentes primarias: '+CCPC+'\nConvocatoria: '+CALL
    pdf.showPage()
prs.save(HERE/'Umbral_presentacion.pptx'); pdf.save()
import fitz
doc=fitz.open(HERE/'Umbral_presentacion.pdf')
for i,page in enumerate(doc): page.get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(HERE/f'QA_{i+1:02}.png')
print('Generadas 10 diapositivas PPTX editables, PDF y 10 vistas QA.')



